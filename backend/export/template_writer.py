"""帳票の穴埋め — ``TemplateWriter`` (f_11 §9.2 / c_16 §4.5.2、段階 B-1b)。

``ExportContent`` は受けない。入力は「``fields`` の語彙 (c_16 §4.5.2) + 値の
dict」で、本文は生成せず既存の様式ファイル (``base``) のプレースホルダへ
値を差し込むだけ (f_11 §9.2 の表: 体裁の継承 (§9.1) とは別経路)。

対象は chat モードのみ。create モードは production_stage が別の経路
(f_08 §3.1.1 の構成テンプレート、または通常の staged/longform 生成) を持ち、
帳票の穴埋めは対象にしない。

``fields`` の語彙・検証・install 時のプレースホルダ突き合わせは
:mod:`backend.free.rag.corpus.templates` が本モジュールの
``parse_field_specs`` / ``scan_placeholders`` / ``required_placeholder_names``
/ ``validate_xlsx_repeat_rows`` を呼んで行う (一方向依存: corpus → export、
逆はしない)。
"""

from __future__ import annotations

import copy
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from backend.export.base import WriteResult
from backend.io.atomic import atomic_write_bytes
from backend.log_config import get_logger

logger = get_logger("export.template_writer")

#: プレースホルダ ``{{name}}`` / ``{{items.品名}}``。
_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
#: セル全体がちょうど 1 個のプレースホルダだけからなるか (xlsx の型付き代入判定)。
_WHOLE_CELL_RE = re.compile(r"^\{\{\s*([^{}]+?)\s*\}\}$")

#: 項目名の字種 (英数字 + ``_``、先頭は英字)。
_NAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")
#: 表の列名は日本語可。プレースホルダ構文と衝突する文字だけ禁止する。
_COLUMN_NAME_RE = re.compile(r"^[^{}.\s]+$")

_FIELD_TYPES = frozenset({"string", "number", "date", "text"})
_REPEAT_KINDS = frozenset({"table_row"})
_COMPUTE_OPS = frozenset({"sum", "sum_product"})


class TemplateFillError(Exception):
    """穴埋めに失敗した (未解決プレースホルダ / 継承元の上書き等)。"""


@dataclass(frozen=True, slots=True)
class ColumnSpec:
    """行繰り返しの 1 列。"""

    name: str
    type: str = "string"


@dataclass(frozen=True, slots=True)
class ComputeSpec:
    """``compute`` (c_16 §4.5.2)。値はコードが計算し、モデルには渡さない。"""

    op: str  # "sum" | "sum_product"
    over: str | None = None
    column: str | None = None
    columns: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class FieldSpec:
    """``fields`` の 1 項目 (検証済み)。"""

    name: str
    type: str = "string"
    required: bool = False
    description: str = ""
    repeat: str | None = None
    columns: tuple[ColumnSpec, ...] = ()
    compute: ComputeSpec | None = None


# ---------------------------------------------------------------------------
# fields の語彙検証
# ---------------------------------------------------------------------------


def _parse_column(raw: Any) -> ColumnSpec | None:
    if not isinstance(raw, dict):
        return None
    name = raw.get("name")
    col_type = raw.get("type", "string")
    if not isinstance(name, str) or not _COLUMN_NAME_RE.match(name.strip()):
        return None
    if col_type not in _FIELD_TYPES:
        return None
    return ColumnSpec(name=name.strip(), type=col_type)


def _parse_compute(raw: Any) -> ComputeSpec | None:
    if not isinstance(raw, dict):
        return None
    op = raw.get("op")
    if op not in _COMPUTE_OPS:
        return None
    over = raw.get("over")
    column = raw.get("column")
    columns_raw = raw.get("columns")
    columns = tuple(columns_raw) if isinstance(columns_raw, list) else ()
    if columns_raw is not None and not all(isinstance(c, str) and c for c in columns):
        return None
    if op == "sum_product":
        if not isinstance(over, str) or not over or len(columns) != 2:
            return None
        return ComputeSpec(op=op, over=over, columns=columns)
    # op == "sum"
    if isinstance(over, str) and over:
        if not isinstance(column, str) or not column:
            return None
        return ComputeSpec(op=op, over=over, column=column)
    if not columns:
        return None
    return ComputeSpec(op=op, columns=columns)


def _parse_field(raw: Any) -> FieldSpec | None:
    if not isinstance(raw, dict):
        return None
    name = raw.get("name")
    if not isinstance(name, str) or not _NAME_RE.match(name):
        return None

    repeat = raw.get("repeat")
    if repeat is not None:
        if repeat not in _REPEAT_KINDS:
            return None
        columns_raw = raw.get("columns")
        if not isinstance(columns_raw, list) or not columns_raw:
            return None
        columns = tuple(_parse_column(c) for c in columns_raw)
        if any(c is None for c in columns):
            return None
        return FieldSpec(name=name, repeat=repeat, columns=columns)  # type: ignore[arg-type]

    compute_raw = raw.get("compute")
    if compute_raw is not None:
        compute = _parse_compute(compute_raw)
        if compute is None:
            return None
        return FieldSpec(
            name=name, type="number",
            description=str(raw.get("description") or ""), compute=compute,
        )

    field_type = raw.get("type", "string")
    if field_type not in _FIELD_TYPES:
        return None
    return FieldSpec(
        name=name, type=field_type, required=bool(raw.get("required", False)),
        description=str(raw.get("description") or ""),
    )


def parse_field_specs(
    raw_fields: "tuple[dict[str, Any], ...] | list[dict[str, Any]]",
) -> tuple[FieldSpec, ...] | None:
    """``fields`` の生 dict 列を検証済みの :class:`FieldSpec` 列にする。

    未知の ``type`` / ``repeat`` / ``compute.op`` が 1 つでもあれば、その
    様式エントリごと無効にする (``None``、c_16 §4.5.2)。名前の重複も無効。
    """
    specs: list[FieldSpec] = []
    seen: set[str] = set()
    for raw in raw_fields:
        spec = _parse_field(raw)
        if spec is None or spec.name in seen:
            return None
        seen.add(spec.name)
        specs.append(spec)
    if not specs:
        return None
    return tuple(specs)


def required_placeholder_names(fields: tuple[FieldSpec, ...]) -> frozenset[str]:
    """``fields`` が要求するプレースホルダ名の集合 (install 時の突き合わせ用)。"""
    names: set[str] = set()
    for f in fields:
        if f.repeat:
            names.update(f"{f.name}.{c.name}" for c in f.columns)
        else:
            names.add(f.name)
    return frozenset(names)


# ---------------------------------------------------------------------------
# プレースホルダの走査 (install 検査と穴埋め本体が共有、run をまたぐ検出込み)
# ---------------------------------------------------------------------------


def _iter_docx_paragraphs(doc: Any):
    """本文段落・表 (入れ子含む)・ヘッダ / フッタの全段落を辿る。"""
    def _table_paragraphs(table: Any):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from _table_paragraphs(nested)

    yield from doc.paragraphs
    for table in doc.tables:
        yield from _table_paragraphs(table)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _scan_docx_placeholders(path: Path) -> set[str]:
    from docx import Document

    doc = Document(str(path))
    found: set[str] = set()
    for paragraph in _iter_docx_paragraphs(doc):
        found.update(m.group(1).strip() for m in _PLACEHOLDER_RE.finditer(paragraph.text))
    return found


def _iter_pptx_paragraphs(prs: Any):
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                yield from shape.text_frame.paragraphs
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield from cell.text_frame.paragraphs


def _scan_pptx_placeholders(path: Path) -> set[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    found: set[str] = set()
    for paragraph in _iter_pptx_paragraphs(prs):
        found.update(m.group(1).strip() for m in _PLACEHOLDER_RE.finditer(paragraph.text))
    return found


def _scan_xlsx_placeholders(path: Path) -> set[str]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False)
    found: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    found.update(
                        m.group(1).strip() for m in _PLACEHOLDER_RE.finditer(cell.value)
                    )
    return found


def scan_placeholders(path: Path) -> set[str]:
    """様式ファイル中のプレースホルダ名の集合を返す (install 検査と穴埋めが共有)。

    開けない / 対応しない拡張子は :class:`TemplateFillError`。呼出側 (install
    検査) はこれを「そのエントリを無効にする」判断材料として使う。
    """
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return _scan_docx_placeholders(path)
        if ext == ".pptx":
            return _scan_pptx_placeholders(path)
        if ext == ".xlsx":
            return _scan_xlsx_placeholders(path)
    except Exception as e:  # noqa: BLE001 - 壊れた/読めないファイルは呼出側が無効化する
        raise TemplateFillError(f"could not scan placeholders in {path}: {e}") from e
    raise TemplateFillError(f"unsupported template base extension: {path.suffix}")


def find_xlsx_repeat_row(ws: Any, field_name: str) -> int | None:
    """``field_name`` の行繰り返しの雛形行 (1-based) を探す (無ければ ``None``)。"""
    needle = f"{{{{{field_name}."
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and needle in cell.value:
                return cell.row
    return None


def validate_xlsx_repeat_rows(path: Path, fields: tuple[FieldSpec, ...]) -> None:
    """雛形行がシートの最終行であること (c_16 §4.5.2 / f_11 §9.2) を検査する。

    違反があれば :class:`TemplateFillError`。install 検査とファイル穴埋めの
    両方から呼ぶ (「install 時と書き込み時の両方で検査する」)。
    """
    repeats = [f for f in fields if f.repeat]
    if not repeats:
        return
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False)
    for f in repeats:
        for ws in wb.worksheets:
            row_idx = find_xlsx_repeat_row(ws, f.name)
            if row_idx is None:
                continue
            if row_idx != ws.max_row:
                raise TemplateFillError(
                    f"xlsx repeat template row for '{f.name}' (row {row_idx}) "
                    f"is not the sheet's last used row ({ws.max_row})",
                )


# ---------------------------------------------------------------------------
# 値の検証・計算
# ---------------------------------------------------------------------------


_LOOSE_DATE_RE = re.compile(
    r"^(\d{4})\s*[年/.\-]\s*(\d{1,2})\s*[月/.\-]\s*(\d{1,2})\s*日?$",
)


def _coerce_number(raw: Any) -> float | None:
    """数値にする。桁区切りと全角は受け、単位や漢数字 (「10万円」) は受けない。

    読めない値を推測で数にしない — ``None`` にして訊き直す側へ倒す。
    ``nan`` / ``inf`` と bool も数として扱わない。
    """
    if isinstance(raw, bool):
        return None
    if isinstance(raw, str):
        raw = unicodedata.normalize("NFKC", raw).strip().replace(",", "")
    try:
        value = float(raw)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def _coerce_date(raw: Any) -> date | None:
    """日付にする。ISO のほか ``2026年9月20日`` / ``2026/9/20`` を受ける。"""
    if isinstance(raw, date):
        return raw
    text = unicodedata.normalize("NFKC", str(raw)).strip()
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    match = _LOOSE_DATE_RE.match(text)
    if match is None:
        return None
    try:
        return date(*(int(g) for g in match.groups()))
    except ValueError:
        return None


def _coerce_cell(field_type: str, raw: Any) -> Any | None:
    if raw is None:
        return None
    if field_type == "number":
        return _coerce_number(raw)
    if field_type == "date":
        return _coerce_date(raw)
    if not isinstance(raw, str):
        return None
    return raw.strip() or None


def _coerce_rows(spec: FieldSpec, raw: Any) -> list[dict[str, Any]] | None:
    """行繰り返しの値を列の型で検査する。読めないセルが 1 つでもあれば ``None``。

    全セルが空の行は捨てる。列に無いキーは落とす。
    """
    if raw is None:
        return []
    if not isinstance(raw, list):
        return None
    rows: list[dict[str, Any]] = []
    for item in raw:
        if not isinstance(item, dict):
            return None
        row: dict[str, Any] = {}
        for column in spec.columns:
            cell = item.get(column.name)
            if cell is None or (isinstance(cell, str) and not cell.strip()):
                row[column.name] = None
                continue
            value = _coerce_cell(column.type, cell)
            if value is None:
                return None
            row[column.name] = value
        if not any(v is not None for v in row.values()):
            continue
        # 中身のある行で数値 / 日付の列が空なら、その表ごと訊き直す。通すと
        # compute が後段で失敗し、「プレースホルダが残った」という分かりにくい
        # 失敗になる (2026-09-20 実機確認: モデルが数量・単価を空で返した)。
        if any(
            row[column.name] is None and column.type in ("number", "date")
            for column in spec.columns
        ):
            return None
        rows.append(row)
    return rows


def validate_values(
    fields: tuple[FieldSpec, ...], values: dict[str, Any],
) -> tuple[bool, list[FieldSpec], dict[str, Any]]:
    """LLM が返した値を型 / ``required`` について検査する (json_schema は形だけ)。

    Returns ``(ok, missing_required, coerced)``。``ok`` が ``False`` の間は
    書かずに ``missing_required`` (description 付き) を訊く。``coerced`` は
    型変換済みの値 (``number`` → ``float``、``date`` → ``date``、行繰り返しは
    列の型で変換した dict の列)。``compute`` はここでは扱わない (別途
    :func:`compute_values`)。

    **値があるのに読めない**もの (「10万円」、数量に「二個」) は ``required`` で
    なくても ``missing`` に入れる — 黙って空欄で書くと、体裁が正しいぶん欠落に
    気付けない。空白だけの文字列は値なしと同じ扱い。
    """
    if not isinstance(values, dict):
        values = {}
    missing: list[FieldSpec] = []
    coerced: dict[str, Any] = {}
    for f in fields:
        if f.compute is not None:
            continue
        raw = values.get(f.name)
        if f.repeat:
            rows = _coerce_rows(f, raw)
            if rows is None or (f.required and not rows):
                missing.append(f)
                rows = []
            coerced[f.name] = rows
            continue
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if f.required:
                missing.append(f)
            continue
        v = _coerce_cell(f.type, raw)
        if v is None:
            missing.append(f)
            continue
        coerced[f.name] = v
    return (not missing, missing, coerced)


def compute_values(
    fields: tuple[FieldSpec, ...], values: dict[str, Any],
) -> dict[str, float]:
    """``compute`` フィールドをコードで計算する (LLM には計算させない)。

    数値化できない入力は **その compute だけ** 書かない
    (未解決のプレースホルダとして最終的にファイル全体の失敗になる、f_11 §9.2)。
    """
    computed: dict[str, float] = {}
    for f in fields:
        c = f.compute
        if c is None:
            continue
        try:
            if c.op == "sum_product":
                rows = values.get(c.over) or []
                col_a, col_b = c.columns
                total = sum(
                    float(row[col_a]) * float(row[col_b]) for row in rows
                )
            elif c.over is not None:
                rows = values.get(c.over) or []
                total = sum(float(row[c.column]) for row in rows)
            else:
                total = sum(float(values[name]) for name in c.columns)
        except (TypeError, ValueError, KeyError):
            logger.warning(
                "template compute '%s' (op=%s) failed on non-numeric input; "
                "leaving unresolved", f.name, c.op,
            )
            continue
        computed[f.name] = total
    return computed


def _render_scalar(field_type: str, value: Any) -> str:
    if value is None:
        return ""
    if field_type == "number":
        v = float(value)
        return str(int(v)) if v.is_integer() else str(v)
    return str(value)


def _scalar_text_map(fields: tuple[FieldSpec, ...], values: dict[str, Any]) -> dict[str, str]:
    """スカラー項目 (行繰り返し以外) の描画済み文字列マップ。

    ``compute`` フィールドで値が無い (計算失敗、:func:`compute_values` が
    キーを立てなかった) 場合は **マップに入れない** — プレースホルダを
    未解決のまま残し、最終的に書込みごと失敗させる (f_11 §9.2: 「値が数値で
    なければその compute は失敗 → 書かない」)。それ以外の任意項目は根拠が
    無ければ空文字 (書込み自体は続行する)。
    """
    text_map: dict[str, str] = {}
    for f in fields:
        if f.repeat:
            continue
        if f.compute is not None and f.name not in values:
            continue
        text_map[f.name] = _render_scalar(f.type, values.get(f.name))
    return text_map


# ---------------------------------------------------------------------------
# run をまたぐプレースホルダ置換 (docx / pptx 共通)
# ---------------------------------------------------------------------------


def _run_index_at(starts: list[int], lengths: list[int], pos: int) -> int:
    for i, (s, length) in enumerate(zip(starts, lengths)):
        if s <= pos < s + length:
            return i
    return max(len(starts) - 1, 0)


def _replace_placeholders_in_paragraph(paragraph: Any, values_map: dict[str, str]) -> bool:
    """段落内の run をまたぐプレースホルダを置換する。

    先頭の run (プレースホルダが始まる run) に置換結果を入れ、残りの該当部分
    を空にする (書式は先頭の run が勝つ、f_11 §9.2)。戻り値は全て解決できたか
    (``False`` なら未解決のプレースホルダが残った)。
    """
    runs = paragraph.runs
    if not runs:
        return True
    full_text = "".join(r.text for r in runs)
    if "{{" not in full_text:
        return True
    matches = list(_PLACEHOLDER_RE.finditer(full_text))
    if not matches:
        return True

    lengths = [len(r.text) for r in runs]
    starts: list[int] = []
    pos = 0
    for length in lengths:
        starts.append(pos)
        pos += length

    edits: dict[int, list[tuple[int, int, str]]] = {}
    resolved_all = True
    for m in matches:
        name = m.group(1).strip()
        if name not in values_map:
            resolved_all = False
            continue
        repl = values_map[name]
        s, e = m.start(), m.end()
        r_start = _run_index_at(starts, lengths, s)
        r_end = _run_index_at(starts, lengths, max(e - 1, s))
        if r_start == r_end:
            local_s = s - starts[r_start]
            local_e = e - starts[r_start]
            edits.setdefault(r_start, []).append((local_s, local_e, repl))
        else:
            local_s = s - starts[r_start]
            edits.setdefault(r_start, []).append((local_s, lengths[r_start], repl))
            for mid in range(r_start + 1, r_end):
                edits.setdefault(mid, []).append((0, lengths[mid], ""))
            local_e = e - starts[r_end]
            edits.setdefault(r_end, []).append((0, local_e, ""))

    for idx, run in enumerate(runs):
        run_edits = edits.get(idx)
        if not run_edits:
            continue
        text = run.text
        for local_s, local_e, repl in sorted(run_edits, key=lambda t: t[0], reverse=True):
            text = text[:local_s] + repl + text[local_e:]
        run.text = text
    return resolved_all


# ---------------------------------------------------------------------------
# .docx
# ---------------------------------------------------------------------------


def _find_docx_repeat_row(doc: Any, field_name: str):
    needle = f"{{{{{field_name}."

    def _search(table: Any):
        for row in table.rows:
            row_text = "".join(p.text for cell in row.cells for p in cell.paragraphs)
            if needle in row_text:
                return row, table
            for cell in row.cells:
                for nested in cell.tables:
                    found = _search(nested)
                    if found is not None:
                        return found
        return None

    for table in doc.tables:
        found = _search(table)
        if found is not None:
            return found
    return None


def _expand_docx_repeat(doc: Any, field: FieldSpec, rows_data: list[dict], scalar_map: dict[str, str]) -> bool:
    from docx.table import _Row

    found = _find_docx_repeat_row(doc, field.name)
    if found is None:
        return True
    template_row, table = found
    tr_element = template_row._tr
    parent = tr_element.getparent()
    if not rows_data:
        parent.remove(tr_element)
        return True

    ok = True
    anchor = tr_element
    for row_values in rows_data:
        new_tr = copy.deepcopy(tr_element)
        anchor.addprevious(new_tr)
        new_row = _Row(new_tr, table)
        row_map = dict(scalar_map)
        for col in field.columns:
            raw = row_values.get(col.name)
            row_map[f"{field.name}.{col.name}"] = _render_scalar(col.type, raw)
        for cell in new_row.cells:
            for paragraph in cell.paragraphs:
                if not _replace_placeholders_in_paragraph(paragraph, row_map):
                    ok = False
            for nested in cell.tables:
                for nrow in nested.rows:
                    for ncell in nrow.cells:
                        for paragraph in ncell.paragraphs:
                            if not _replace_placeholders_in_paragraph(paragraph, row_map):
                                ok = False
    parent.remove(tr_element)
    return ok


def _fill_docx(base_path: Path, fields: tuple[FieldSpec, ...], values: dict[str, Any]) -> bytes:
    import io

    from docx import Document

    doc = Document(str(base_path))
    scalar_map = _scalar_text_map(fields, values)
    ok = True
    for f in fields:
        if f.repeat:
            if not _expand_docx_repeat(doc, f, values.get(f.name) or [], scalar_map):
                ok = False
    for paragraph in _iter_docx_paragraphs(doc):
        if not _replace_placeholders_in_paragraph(paragraph, scalar_map):
            ok = False
    if not ok:
        raise TemplateFillError("unresolved placeholder(s) remained in the docx template")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# .pptx
# ---------------------------------------------------------------------------


def _fill_pptx(base_path: Path, fields: tuple[FieldSpec, ...], values: dict[str, Any]) -> bytes:
    import io

    from pptx import Presentation

    prs = Presentation(str(base_path))
    scalar_map = _scalar_text_map(fields, values)
    ok = True
    for paragraph in _iter_pptx_paragraphs(prs):
        if not _replace_placeholders_in_paragraph(paragraph, scalar_map):
            ok = False
    if not ok:
        raise TemplateFillError("unresolved placeholder(s) remained in the pptx template")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# .xlsx
# ---------------------------------------------------------------------------


def _replace_xlsx_text(text: str, resolver) -> tuple[str, bool]:
    ok = True

    def _repl(m: re.Match) -> str:
        nonlocal ok
        name = m.group(1).strip()
        value = resolver(name)
        if value is None:
            ok = False
            return m.group(0)
        return value

    return _PLACEHOLDER_RE.sub(_repl, text), ok


def _fill_xlsx_cell(cell: Any, field_by_name: dict[str, FieldSpec], values: dict[str, Any], scalar_map: dict[str, str]) -> bool:
    text = cell.value
    if not isinstance(text, str) or "{{" not in text:
        return True
    whole = _WHOLE_CELL_RE.match(text.strip())
    if whole is not None:
        name = whole.group(1).strip()
        f = field_by_name.get(name)
        if f is not None and f.type == "number" and values.get(name) is not None:
            cell.value = values[name]
            return True
        if f is not None and f.type == "date" and values.get(name) is not None:
            cell.value = values[name]
            return True
        if name in scalar_map:
            cell.value = scalar_map[name]
            return True
        return False
    new_text, ok = _replace_xlsx_text(text, lambda name: scalar_map.get(name))
    cell.value = new_text
    return ok


def _expand_xlsx_repeat(ws: Any, field: FieldSpec, rows_data: list[dict]) -> bool:
    from copy import copy as _copy_style

    row_idx = find_xlsx_repeat_row(ws, field.name)
    if row_idx is None:
        return True
    if not rows_data:
        ws.delete_rows(row_idx, 1)
        return True
    if row_idx != ws.max_row:
        raise TemplateFillError(
            f"xlsx repeat template row for '{field.name}' is not the sheet's "
            f"last used row",
        )
    max_col = ws.max_column
    col_by_name = {c.name: c for c in field.columns}
    template_texts = [ws.cell(row=row_idx, column=col).value for col in range(1, max_col + 1)]

    def _fill_row(idx: int, row_values: dict) -> bool:
        ok = True
        prefix = f"{field.name}."
        for col in range(1, max_col + 1):
            cell = ws.cell(row=idx, column=col)
            text = cell.value
            if not isinstance(text, str) or "{{" not in text:
                continue
            whole = _WHOLE_CELL_RE.match(text.strip())
            if whole is not None and whole.group(1).strip().startswith(prefix):
                colname = whole.group(1).strip()[len(prefix):]
                colspec = col_by_name.get(colname)
                raw = row_values.get(colname)
                if colspec is not None and colspec.type == "number":
                    v = _coerce_number(raw)
                    if v is None:
                        ok = False
                    else:
                        cell.value = v
                    continue
                if colspec is not None and colspec.type == "date":
                    v = _coerce_date(raw)
                    if v is None:
                        ok = False
                    else:
                        cell.value = v
                    continue
                cell.value = "" if raw is None else str(raw)
                continue

            def _resolve(name: str, rv=row_values) -> str | None:
                if name.startswith(prefix):
                    v = rv.get(name[len(prefix):])
                    return "" if v is None else str(v)
                return None

            new_text, row_ok = _replace_xlsx_text(text, _resolve)
            cell.value = new_text
            ok = ok and row_ok
        return ok

    ok = True
    for i, row_values in enumerate(rows_data):
        idx = row_idx + i
        if i > 0:
            for col in range(1, max_col + 1):
                src = ws.cell(row=row_idx, column=col)
                dst = ws.cell(row=idx, column=col)
                dst.value = template_texts[col - 1]
                dst.font = _copy_style(src.font)
                dst.fill = _copy_style(src.fill)
                dst.border = _copy_style(src.border)
                dst.alignment = _copy_style(src.alignment)
                dst.number_format = src.number_format
        if not _fill_row(idx, row_values):
            ok = False
    return ok


def _fill_xlsx(base_path: Path, fields: tuple[FieldSpec, ...], values: dict[str, Any]) -> bytes:
    import io

    from openpyxl import load_workbook

    wb = load_workbook(str(base_path), data_only=False)
    field_by_name = {f.name: f for f in fields if not f.repeat}
    scalar_map = _scalar_text_map(fields, values)
    ok = True
    for ws in wb.worksheets:
        for f in fields:
            if f.repeat:
                if not _expand_xlsx_repeat(ws, f, values.get(f.name) or []):
                    ok = False
        for row in ws.iter_rows():
            for cell in row:
                if not _fill_xlsx_cell(cell, field_by_name, values, scalar_map):
                    ok = False
    if not ok:
        raise TemplateFillError("unresolved placeholder(s) remained in the xlsx template")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 公開 API
# ---------------------------------------------------------------------------


_FILLERS = {".docx": _fill_docx, ".pptx": _fill_pptx, ".xlsx": _fill_xlsx}


def fill_template(
    base_path: Path,
    fields: "tuple[FieldSpec, ...] | tuple[dict[str, Any], ...]",
    values: dict[str, Any],
    out_path: Path,
    *,
    provenance: str | None = None,
) -> WriteResult:
    """様式 (``base_path``) のプレースホルダへ ``values`` を差し込んで書き出す。

    ``fields`` は :func:`parse_field_specs` 済みでも生 dict 列でもよい (呼出側
    の手間を減らすため、ここで一度だけ検証する)。未解決のプレースホルダが
    1 つでも残れば書かずに :class:`TemplateFillError`。継承元 (``base_path``)
    と同じ ``out_path`` へは書かない。``provenance`` は
    ``<package_id>@<version>:<entry_id>`` (f_11 §9.2)。
    """
    base_path = Path(base_path)
    out_path = Path(out_path)
    if out_path.resolve() == base_path.resolve():
        raise TemplateFillError("refusing to overwrite the template base file")

    field_specs = fields if (fields and isinstance(fields[0], FieldSpec)) else parse_field_specs(fields)  # type: ignore[arg-type]
    if not field_specs:
        raise TemplateFillError("invalid or empty field specification")

    ext = base_path.suffix.lower()
    filler = _FILLERS.get(ext)
    if filler is None:
        raise TemplateFillError(f"unsupported template base extension: {ext}")

    data = filler(base_path, field_specs, values)
    atomic_write_bytes(out_path, data)
    metadata: dict[str, Any] = {"template_applied": True}
    if provenance:
        metadata["template"] = provenance
    return WriteResult(path=out_path, size_bytes=len(data), metadata=metadata)


__all__ = [
    "ColumnSpec",
    "ComputeSpec",
    "FieldSpec",
    "TemplateFillError",
    "compute_values",
    "fill_template",
    "find_xlsx_repeat_row",
    "parse_field_specs",
    "required_placeholder_names",
    "scan_placeholders",
    "validate_values",
    "validate_xlsx_repeat_rows",
]
